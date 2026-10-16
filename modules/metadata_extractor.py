"""
metadata_extractor.py — File metadata forensics (deep edition).

Extracts:
  · EXIF/IPTC/XMP from images (PIL, pyexiv2 when available)
  · Full GPS decode + optional reverse geocoding via Nominatim
  · PDF metadata + XMP (pypdf, pikepdf when available)
  · Office documents (docx / openpyxl / python-pptx / zipfile fallback)
  · Audio tags (mutagen: MP3, FLAC, MP4, OGG, WAV, WMA, AIFF)
  · Video metadata (ffprobe — format + streams)
  · Archive contents (zip / tar / 7z / rar via zipfile / tarfile / py7zr)
  · Executables (PE headers, ELF, Mach-O, scripts)
  · Cryptographic hashes (MD5, SHA1, SHA256, SHA512)
  · Shannon entropy + packed/encrypted heuristic
  · MIME detection via python-magic (fallback to extension)
  · LSB + chi-square steganography detection (images)
  · Timestamp anomaly detection (fs vs metadata drift)
  · Sensitive-data scan across all extracted fields
  · Risk scoring 0–100 with plain-English reasons

Also provides:
  · Metadata scrubber for images, PDFs, Office docs, archives
  · Batch directory walker with parallel parsing
  · Scan history per file with diff
  · Exports: JSON, CSV, standalone HTML report
"""

import csv
import hashlib
import html
import io
import json
import math
import mimetypes
import os
import re
import struct
import subprocess
import tarfile
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path

import requests
from rich.console import Console
from rich.panel import Panel
from rich.table import Table
from rich.theme import Theme
from tqdm import tqdm

console = Console(theme=Theme({
    "g": "grey70", "dg": "grey50", "bg": "grey93", "bd": "grey37",
    "red": "red", "green": "green", "yellow": "yellow",
}))

BAR_FMT = "\033[38;5;245m  {desc} |{bar}| {percentage:3.0f}%\033[0m"
OUT_DIR = Path.home() / ".ostinttol" / "exports"
OUT_DIR.mkdir(parents=True, exist_ok=True)
DATA_DIR = Path.home() / ".ostinttol" / "metadata_extractor"
DATA_DIR.mkdir(parents=True, exist_ok=True)
HIST_DIR = DATA_DIR / "history"
HIST_DIR.mkdir(parents=True, exist_ok=True)
CONFIG_FILE = Path("config") / "config.json"

UA = "Mozilla/5.0 (compatible; ostinttol)"


# ============================================================
# Optional imports
# ============================================================
try:
    from PIL import Image, ImageChops, ImageFilter, ImageStat
    from PIL.ExifTags import TAGS, GPSTAGS
    HAS_PIL = True
except ImportError:
    HAS_PIL = False

try:
    from pypdf import PdfReader
    HAS_PYPDF = True
except ImportError:
    try:
        from PyPDF2 import PdfReader
        HAS_PYPDF = True
    except ImportError:
        HAS_PYPDF = False

try:
    import pikepdf
    HAS_PIKEPDF = True
except ImportError:
    HAS_PIKEPDF = False

try:
    from mutagen import File as MutagenFile
    HAS_MUTAGEN = True
except ImportError:
    HAS_MUTAGEN = False

try:
    import magic as python_magic
    HAS_MAGIC = True
except ImportError:
    HAS_MAGIC = False

try:
    import py7zr
    HAS_PY7ZR = True
except ImportError:
    HAS_PY7ZR = False

try:
    import rarfile
    HAS_RARFILE = True
except ImportError:
    HAS_RARFILE = False

try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    from pptx import Presentation
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False


# ============================================================
# Config
# ============================================================
def _load_config():
    if CONFIG_FILE.exists():
        try:
            return json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
        except Exception:
            pass
    return {}


# ============================================================
# Extensions
# ============================================================
IMAGE_EXTS = {
    ".jpg", ".jpeg", ".png", ".tif", ".tiff", ".webp", ".heic", ".heif",
    ".bmp", ".gif", ".avif", ".jxl", ".ico",
    ".cr2", ".cr3", ".nef", ".nrw", ".arw", ".srf", ".sr2", ".dng",
    ".orf", ".rw2", ".raf", ".pef", ".srw", ".x3f", ".3fr", ".fff",
}
PDF_EXTS = {".pdf"}
OFFICE_EXTS = {
    ".docx", ".docm", ".xlsx", ".xlsm", ".pptx", ".pptm",
    ".odt", ".ods", ".odp", ".odg",
}
ARCHIVE_EXTS = {
    ".zip", ".jar", ".apk", ".epub", ".war", ".ear",
    ".tar", ".gz", ".tgz", ".bz2", ".xz", ".7z", ".rar", ".iso",
}
AUDIO_EXTS = {
    ".mp3", ".flac", ".m4a", ".aac", ".ogg", ".oga", ".opus",
    ".wav", ".wma", ".aiff", ".aif", ".ape", ".wv", ".mpc",
    ".tta", ".dsf", ".dff", ".mka", ".ac3",
}
VIDEO_EXTS = {
    ".mp4", ".m4v", ".mov", ".avi", ".mkv", ".webm", ".flv", ".wmv",
    ".mpg", ".mpeg", ".m2ts", ".ts", ".3gp", ".ogv", ".vob",
}
EXECUTABLE_EXTS = {
    ".exe", ".dll", ".sys", ".scr", ".ocx", ".cpl", ".so", ".dylib",
    ".elf", ".bin", ".com", ".msi", ".mach-o",
    ".sh", ".bash", ".zsh", ".ps1", ".bat", ".cmd", ".vbs",
    ".py", ".rb", ".pl", ".js", ".php",
}
TEXT_EXTS = {
    ".txt", ".md", ".log", ".ini", ".cfg", ".conf", ".json",
    ".xml", ".yaml", ".yml", ".toml", ".csv", ".tsv", ".html",
    ".htm", ".rtf", ".tex",
}


# ============================================================
# Hashing
# ============================================================
def hash_file(path: Path, chunk_size: int = 1 << 20):
    """Compute MD5 / SHA1 / SHA256 / SHA512 in a single pass."""
    md5 = hashlib.md5()
    sha1 = hashlib.sha1()
    sha256 = hashlib.sha256()
    sha512 = hashlib.sha512()
    size = 0
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(chunk_size), b""):
            md5.update(chunk)
            sha1.update(chunk)
            sha256.update(chunk)
            sha512.update(chunk)
            size += len(chunk)
    return {
        "md5": md5.hexdigest(),
        "sha1": sha1.hexdigest(),
        "sha256": sha256.hexdigest(),
        "sha512": sha512.hexdigest(),
        "size_bytes": size,
    }


# ============================================================
# Entropy
# ============================================================
def entropy_of_bytes(data: bytes, sample_size: int = 1 << 20) -> float:
    """Shannon entropy of a byte sample (0-8 bits per byte)."""
    if not data:
        return 0.0
    sample = data[:sample_size]
    freq = [0] * 256
    for b in sample:
        freq[b] += 1
    length = len(sample)
    ent = 0.0
    for count in freq:
        if count:
            p = count / length
            ent -= p * math.log2(p)
    return round(ent, 4)


def entropy_file(path: Path, sample_size: int = 1 << 20) -> float:
    try:
        with path.open("rb") as fh:
            return entropy_of_bytes(fh.read(sample_size), sample_size)
    except Exception:
        return 0.0


def entropy_label(ent: float) -> str:
    if ent < 3.0:
        return "very low (likely text or structure)"
    if ent < 5.0:
        return "low (mostly text or uncompressed)"
    if ent < 6.5:
        return "medium (mixed content)"
    if ent < 7.5:
        return "high (likely compressed)"
    return "very high (likely encrypted or compressed)"


# ============================================================
# MIME detection
# ============================================================
def detect_mime(path: Path):
    """Return MIME via python-magic when available, else mimetypes."""
    if HAS_MAGIC:
        try:
            m = python_magic.Magic(mime=True)
            return m.from_file(str(path))
        except Exception:
            pass
    guess, _ = mimetypes.guess_type(str(path))
    return guess or "application/octet-stream"


def detect_description(path: Path):
    """Human-readable file description via python-magic."""
    if not HAS_MAGIC:
        return None
    try:
        m = python_magic.Magic(mime=False)
        return m.from_file(str(path))
    except Exception:
        return None


# ============================================================
# Classification
# ============================================================
def classify(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in IMAGE_EXTS:
        return "image"
    if ext in PDF_EXTS:
        return "pdf"
    if ext in OFFICE_EXTS:
        return "office"
    if ext in ARCHIVE_EXTS:
        return "archive"
    if ext in AUDIO_EXTS:
        return "audio"
    if ext in VIDEO_EXTS:
        return "video"
    if ext in EXECUTABLE_EXTS:
        return "executable"
    if ext in TEXT_EXTS:
        return "text"
    return "unknown"


# ============================================================
# GPS helpers
# ============================================================
def _rational(v):
    try:
        if isinstance(v, tuple) and len(v) == 2:
            return v[0] / v[1] if v[1] else 0.0
        return float(v)
    except Exception:
        return None


def _dms_to_decimal(dms_val, ref):
    try:
        d = _rational(dms_val[0])
        m = _rational(dms_val[1])
        s = _rational(dms_val[2])
        if None in (d, m, s):
            return None
        val = d + m / 60 + s / 3600
        if str(ref).upper() in ("S", "W"):
            val = -val
        return round(val, 7)
    except Exception:
        return None


def _gps_decimal(gps):
    if not gps:
        return None
    try:
        lat = _dms_to_decimal(gps.get("GPSLatitude"),
                              gps.get("GPSLatitudeRef", "N"))
        lon = _dms_to_decimal(gps.get("GPSLongitude"),
                              gps.get("GPSLongitudeRef", "E"))
        alt = _rational(gps.get("GPSAltitude"))
        has = lat is not None and lon is not None
        out = {
            "lat": lat,
            "lon": lon,
            "altitude_m": round(alt, 2) if alt is not None else None,
            "osm": (f"https://www.openstreetmap.org/?mlat={lat}&mlon={lon}"
                    f"#map=17/{lat}/{lon}") if has else None,
            "google": (f"https://www.google.com/maps?q={lat},{lon}")
                      if has else None,
        }
        # timestamp and direction when present
        ts = gps.get("GPSTimeStamp")
        ds = gps.get("GPSDateStamp")
        if ts or ds:
            try:
                out["gps_timestamp"] = f"{ds or ''} {ts or ''}".strip()
            except Exception:
                pass
        return out
    except Exception:
        return None


def reverse_geocode(lat, lon, timeout=10):
    """Nominatim reverse geocode — OSM public endpoint (1 req/sec limit)."""
    if lat is None or lon is None:
        return None
    try:
        r = requests.get(
            "https://nominatim.openstreetmap.org/reverse",
            params={"lat": lat, "lon": lon, "format": "jsonv2"},
            headers={"User-Agent": "ostinttol/1.0 (metadata extractor)"},
            timeout=timeout,
        )
        if r.status_code == 200:
            d = r.json()
            return {
                "display_name": d.get("display_name"),
                "address": d.get("address"),
                "osm_type": d.get("osm_type"),
                "osm_id": d.get("osm_id"),
                "place_id": d.get("place_id"),
            }
    except Exception:
        pass
    return None


# ============================================================
# Image metadata
# ============================================================
def _image_meta(path: Path, deep=True):
    if not HAS_PIL:
        return {"error": "Pillow not installed"}
    try:
        img = Image.open(path)
    except Exception as e:
        return {"error": f"open failed: {e}"}

    info = {
        "format": img.format,
        "mode": img.mode,
        "width": img.width,
        "height": img.height,
        "size": f"{img.width}x{img.height}",
        "aspect_ratio": round(img.width / img.height, 4) if img.height else None,
        "megapixels": round(img.width * img.height / 1e6, 2),
    }

    try:
        exif = img.getexif() or {}
    except Exception:
        exif = {}

    tags = {}
    for tid, val in exif.items():
        name = TAGS.get(tid, tid)
        if name == "GPSInfo":
            gps = {}
            for k, v in (val or {}).items():
                gps[GPSTAGS.get(k, k)] = v
            tags["GPSInfo"] = gps
        else:
            tags[name] = val if isinstance(val, (str, int, float)) else str(val)

    # sub-IFD exif block
    try:
        sub = exif.get_ifd(0x8769) or {}
        for tid, val in sub.items():
            name = TAGS.get(tid, tid)
            tags[name] = val if isinstance(val, (str, int, float)) else str(val)
    except Exception:
        pass

    # GPS IFD
    try:
        gps_ifd = exif.get_ifd(0x8825) or {}
        if gps_ifd:
            gps_named = {}
            for k, v in gps_ifd.items():
                gps_named[GPSTAGS.get(k, k)] = v
            tags["GPSInfo"] = gps_named
    except Exception:
        pass

    info["exif"] = tags
    info["gps"] = _gps_decimal(tags.get("GPSInfo"))
    info["exif_tag_count"] = len(tags)

    # PNG text chunks
    if img.format == "PNG":
        try:
            png_text = {}
            for k, v in (img.text or {}).items():
                png_text[k] = v[:200] if isinstance(v, str) else str(v)
            info["png_text"] = png_text
        except Exception:
            pass
        try:
            info["png_info"] = dict(img.info or {})
        except Exception:
            pass

    # palette / transparency
    try:
        if img.mode == "P":
            info["palette_size"] = len(img.getpalette() or []) // 3
    except Exception:
        pass

    # deep analysis
    if deep:
        info["blur_score"] = _blur_score(img)
        info["dominant_hex"] = _dominant_hex(img)
        info["brightness"], info["contrast"] = _brightness_contrast(img)
        info["stego"] = stego_analysis(img)

    return info


def _blur_score(img):
    try:
        gray = img.convert("L")
        lap = gray.filter(ImageFilter.Kernel(
            (3, 3), [0, 1, 0, 1, -4, 1, 0, 1, 0], scale=1, offset=128))
        st = ImageStat.Stat(lap)
        return round(st.var[0] if st.var else 0, 2)
    except Exception:
        return None


def _dominant_hex(img):
    try:
        small = img.convert("RGB").resize((64, 64))
        px = list(small.getdata())
        n = len(px) or 1
        r = sum(p[0] for p in px) // n
        g = sum(p[1] for p in px) // n
        b = sum(p[2] for p in px) // n
        return f"#{r:02x}{g:02x}{b:02x}"
    except Exception:
        return None


def _brightness_contrast(img):
    try:
        gray = img.convert("L")
        st = ImageStat.Stat(gray)
        bright = round(st.mean[0], 2) if st.mean else 0
        contrast = round(st.stddev[0], 2) if st.stddev else 0
        return bright, contrast
    except Exception:
        return None, None


# ============================================================
# Steganography — LSB + chi-square
# ============================================================
def stego_analysis(img):
    """LSB plane statistics + chi-square (Westfeld pair frequency)."""
    try:
        rgb = img.convert("RGB")
        px = list(rgb.getdata())
        n = len(px)
        if n == 0:
            return None
        sample = px[:min(n, 400000)]
        sz = len(sample)
        if sz == 0:
            return None

        lsb = [0, 0, 0]
        for p in sample:
            for i in range(3):
                lsb[i] += p[i] & 1
        ratio = [round(x / sz, 6) for x in lsb]

        def chi_sq(ch):
            pairs = {}
            for p in sample:
                v = p[ch]
                even = v if v % 2 == 0 else v - 1
                pairs.setdefault(even, [0, 0])
                pairs[even][v % 2] += 1
            chi = 0.0
            for even, (e, o) in pairs.items():
                exp = (e + o) / 2.0
                if exp > 0:
                    chi += ((e - exp) ** 2 + (o - exp) ** 2) / exp
            return chi

        chi = [round(chi_sq(i), 4) for i in range(3)]
        chi_avg = round(sum(chi) / 3.0, 4)

        suspicious = []
        for i, ch in enumerate("RGB"):
            if abs(ratio[i] - 0.5) < 0.02:
                suspicious.append(f"LSB {ch} near 50/50")
        if chi_avg < 50:
            suspicious.append(f"low chi-square ({chi_avg})")

        likelihood = ("high" if len(suspicious) >= 2
                      else "medium" if suspicious else "low")

        return {
            "lsb_ratio_r": ratio[0],
            "lsb_ratio_g": ratio[1],
            "lsb_ratio_b": ratio[2],
            "chi_square_r": chi[0],
            "chi_square_g": chi[1],
            "chi_square_b": chi[2],
            "chi_square_avg": chi_avg,
            "suspicious": suspicious,
            "likelihood": likelihood,
        }
    except Exception:
        return None


# ============================================================
# PDF metadata
# ============================================================
def _pdf_meta(path: Path):
    out = {}
    if HAS_PYPDF:
        try:
            r = PdfReader(str(path))
            out["pages"] = len(r.pages)
            out["encrypted"] = r.is_encrypted
            meta = dict(r.metadata or {})
            out["metadata"] = {str(k): str(v) for k, v in meta.items()}
            try:
                xmp = r.xmp_metadata
                if xmp is not None:
                    fields = {}
                    for k in dir(xmp):
                        if k.startswith("_"):
                            continue
                        try:
                            v = getattr(xmp, k)
                        except Exception:
                            continue
                        if callable(v):
                            continue
                        if v is None:
                            continue
                        fields[k] = str(v)[:500]
                    out["xmp"] = fields
            except Exception:
                pass
        except Exception as e:
            out["pypdf_error"] = str(e)
    else:
        out["pypdf_error"] = "pypdf not installed"

    if HAS_PIKEPDF:
        try:
            with pikepdf.open(str(path)) as pdf:
                docinfo = {}
                if pdf.docinfo:
                    for k, v in pdf.docinfo.items():
                        try:
                            docinfo[str(k)] = str(v)
                        except Exception:
                            docinfo[str(k)] = "<unreadable>"
                out["docinfo"] = docinfo
                try:
                    with pdf.open_metadata() as meta:
                        xmp = {k: str(meta[k])[:500] for k in meta.keys()}
                        if xmp:
                            out["pikepdf_xmp"] = xmp
                except Exception:
                    pass
                out["pdf_version"] = pdf.pdf_version
                out["pages_pikepdf"] = len(pdf.pages)
                try:
                    attachments = []
                    for name, obj in pdf.attachments.items():
                        attachments.append(str(name))
                    if attachments:
                        out["attachments"] = attachments
                except Exception:
                    pass
        except Exception as e:
            out["pikepdf_error"] = str(e)

    return out


# ============================================================
# Office documents
# ============================================================
def _office_meta(path: Path):
    out = {"type": path.suffix.lower()}

    # try native library first for structural detail
    ext = path.suffix.lower()
    if ext in (".docx", ".docm") and HAS_DOCX:
        try:
            d = docx.Document(str(path))
            cp = d.core_properties
            out["docx_core"] = {
                "title": cp.title,
                "subject": cp.subject,
                "author": cp.author,
                "last_modified_by": cp.last_modified_by,
                "created": str(cp.created),
                "modified": str(cp.modified),
                "category": cp.category,
                "comments": cp.comments,
                "keywords": cp.keywords,
                "language": cp.language,
                "revision": cp.revision,
                "content_status": cp.content_status,
                "identifier": cp.identifier,
                "version": cp.version,
            }
            out["paragraphs"] = len(d.paragraphs)
            out["tables"] = len(d.tables)
        except Exception:
            pass

    if ext in (".xlsx", ".xlsm") and HAS_OPENPYXL:
        try:
            wb = openpyxl.load_workbook(str(path), read_only=True,
                                        data_only=False)
            props = wb.properties
            out["xlsx_core"] = {
                "title": props.title,
                "creator": props.creator,
                "lastModifiedBy": props.lastModifiedBy,
                "created": str(props.created),
                "modified": str(props.modified),
                "category": props.category,
                "description": props.description,
                "keywords": props.keywords,
                "subject": props.subject,
                "revision": props.revision,
            }
            out["sheets"] = wb.sheetnames
            wb.close()
        except Exception:
            pass

    if ext in (".pptx", ".pptm") and HAS_PPTX:
        try:
            pres = Presentation(str(path))
            cp = pres.core_properties
            out["pptx_core"] = {
                "title": cp.title,
                "author": cp.author,
                "last_modified_by": cp.last_modified_by,
                "created": str(cp.created),
                "modified": str(cp.modified),
                "subject": cp.subject,
                "keywords": cp.keywords,
                "category": cp.category,
                "comments": cp.comments,
            }
            out["slides"] = len(pres.slides)
        except Exception:
            pass

    # always walk the zip structure
    try:
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
            out["entries"] = len(names)
            for candidate in ("docProps/core.xml", "docProps/app.xml",
                              "docProps/custom.xml", "meta.xml"):
                if candidate in names:
                    with z.open(candidate) as fh:
                        raw = fh.read().decode("utf-8", errors="ignore")
                    pairs = re.findall(
                        r"<([a-zA-Z]+:)?([a-zA-Z]+)[^>]*>([^<]+)<", raw)
                    extracted = {}
                    for _, tag, val in pairs:
                        val = val.strip()
                        if val:
                            extracted[tag] = val[:500]
                    out[candidate] = extracted
    except Exception as e:
        out["zip_error"] = str(e)

    return out


# ============================================================
# Archives
# ============================================================
def _archive_meta(path: Path):
    ext = path.suffix.lower()
    out = {"archive_type": ext}
    if ext in (".zip", ".jar", ".apk", ".epub", ".war", ".ear"):
        try:
            with zipfile.ZipFile(path) as z:
                entries = z.infolist()
                out["entries"] = len(entries)
                out["total_size"] = sum(e.file_size for e in entries)
                out["compressed_size"] = sum(e.compress_size for e in entries)
                try:
                    out["comment"] = z.comment.decode(errors="ignore") or None
                except Exception:
                    out["comment"] = None
                out["sample"] = [e.filename for e in entries[:30]]
                # per-entry timestamps
                ts = {}
                for e in entries[:30]:
                    if e.date_time and e.date_time[0] >= 1980:
                        try:
                            ts[e.filename] = datetime(*e.date_time).isoformat(
                                timespec="seconds")
                        except Exception:
                            continue
                if ts:
                    out["entry_timestamps"] = ts
                # is there embedded metadata?
                for candidate in ("META-INF/MANIFEST.MF",
                                  "META-INF/maven/pom.properties",
                                  "AndroidManifest.xml",
                                  "META-INF/container.xml"):
                    if candidate in z.namelist():
                        with z.open(candidate) as fh:
                            out[candidate] = fh.read(2048).decode(
                                "utf-8", errors="ignore")
        except Exception as e:
            out["zip_error"] = str(e)

    elif ext in (".tar", ".gz", ".tgz", ".bz2", ".xz"):
        try:
            mode = "r:*"
            with tarfile.open(str(path), mode) as t:
                members = t.getmembers()
                out["entries"] = len(members)
                out["total_size"] = sum(m.size for m in members)
                out["sample"] = [m.name for m in members[:30]]
                ts = {}
                for m in members[:30]:
                    if m.mtime:
                        try:
                            ts[m.name] = datetime.fromtimestamp(
                                m.mtime).isoformat(timespec="seconds")
                        except Exception:
                            continue
                if ts:
                    out["entry_timestamps"] = ts
        except Exception as e:
            out["tar_error"] = str(e)

    elif ext == ".7z" and HAS_PY7ZR:
        try:
            with py7zr.SevenZipFile(str(path), mode="r") as z:
                names = z.getnames()
                out["entries"] = len(names)
                out["sample"] = names[:30]
        except Exception as e:
            out["7z_error"] = str(e)

    elif ext == ".rar" and HAS_RARFILE:
        try:
            with rarfile.RarFile(str(path)) as r:
                entries = r.infolist()
                out["entries"] = len(entries)
                out["sample"] = [e.filename for e in entries[:30]]
        except Exception as e:
            out["rar_error"] = str(e)

    return out


# ============================================================
# Audio
# ============================================================
def _audio_meta(path: Path):
    if not HAS_MUTAGEN:
        return {"error": "mutagen not installed"}
    try:
        m = MutagenFile(str(path))
    except Exception as e:
        return {"error": f"mutagen open failed: {e}"}
    if m is None:
        return {"error": "mutagen could not identify the file"}

    out = {"mutagen_type": type(m).__name__}
    try:
        if m.info:
            info = m.info
            for attr in ("length", "bitrate", "sample_rate", "channels",
                         "bits_per_sample", "codec", "codec_description"):
                if hasattr(info, attr):
                    try:
                        out[attr] = getattr(info, attr)
                    except Exception:
                        continue
    except Exception:
        pass

    try:
        if m.tags:
            tags = {}
            for k, v in m.tags.items():
                try:
                    if isinstance(v, list):
                        v = [str(x) for x in v]
                    tags[str(k)] = str(v)[:500]
                except Exception:
                    tags[str(k)] = "<unreadable>"
            out["tags"] = tags
            out["tag_count"] = len(tags)
    except Exception:
        pass

    return out


# ============================================================
# Video (ffprobe)
# ============================================================
def _ffprobe_available():
    try:
        subprocess.run(["ffprobe", "-version"],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                       timeout=5, check=True)
        return True
    except Exception:
        return False


def _video_meta(path: Path):
    if not _ffprobe_available():
        return {"error": "ffprobe not installed (install ffmpeg)"}
    try:
        proc = subprocess.run(
            ["ffprobe", "-v", "quiet", "-print_format", "json",
             "-show_format", "-show_streams", str(path)],
            capture_output=True, text=True, timeout=30, check=True,
        )
        d = json.loads(proc.stdout)
    except Exception as e:
        return {"error": f"ffprobe failed: {e}"}

    out = {}
    fmt = d.get("format") or {}
    out["format_name"] = fmt.get("format_name")
    out["format_long_name"] = fmt.get("format_long_name")
    out["duration"] = fmt.get("duration")
    out["size"] = fmt.get("size")
    out["bit_rate"] = fmt.get("bit_rate")
    if fmt.get("tags"):
        out["format_tags"] = {k: str(v)[:500] for k, v in fmt["tags"].items()}

    streams = []
    for s in (d.get("streams") or []):
        entry = {
            "index": s.get("index"),
            "codec_type": s.get("codec_type"),
            "codec_name": s.get("codec_name"),
            "codec_long_name": s.get("codec_long_name"),
            "profile": s.get("profile"),
            "width": s.get("width"),
            "height": s.get("height"),
            "sample_rate": s.get("sample_rate"),
            "channels": s.get("channels"),
            "channel_layout": s.get("channel_layout"),
            "bit_rate": s.get("bit_rate"),
            "duration": s.get("duration"),
            "nb_frames": s.get("nb_frames"),
            "r_frame_rate": s.get("r_frame_rate"),
            "pix_fmt": s.get("pix_fmt"),
        }
        if s.get("tags"):
            entry["tags"] = {k: str(v)[:500] for k, v in s["tags"].items()}
        streams.append({k: v for k, v in entry.items() if v is not None})
    out["streams"] = streams

    return out


# ============================================================
# Executables
# ============================================================
def _executable_meta(path: Path):
    out = {}
    try:
        with path.open("rb") as fh:
            head = fh.read(512)
    except Exception as e:
        return {"error": str(e)}

    # ELF
    if head[:4] == b"\x7fELF":
        out["format"] = "ELF"
        try:
            ei_class = head[4]
            ei_data = head[5]
            out["elf_class"] = "64-bit" if ei_class == 2 else "32-bit"
            out["elf_endian"] = "little" if ei_data == 1 else "big"
            e_type = struct.unpack("<H" if ei_data == 1 else ">H", head[16:18])[0]
            types = {1: "relocatable", 2: "executable", 3: "shared object",
                     4: "core"}
            out["elf_type"] = types.get(e_type, f"unknown ({e_type})")
            machine = struct.unpack(
                "<H" if ei_data == 1 else ">H", head[18:20])[0]
            machines = {0x03: "x86", 0x3E: "x86-64", 0x28: "ARM",
                        0xB7: "AArch64", 0xF3: "RISC-V"}
            out["elf_machine"] = machines.get(machine, hex(machine))
        except Exception:
            pass

    # PE
    elif head[:2] == b"MZ":
        out["format"] = "PE"
        try:
            pe_offset = struct.unpack("<I", head[0x3C:0x40])[0]
            with path.open("rb") as fh:
                fh.seek(pe_offset)
                sig = fh.read(4)
                if sig == b"PE\x00\x00":
                    machine = struct.unpack("<H", fh.read(2))[0]
                    machines = {0x014c: "x86", 0x8664: "x86-64",
                                0x01c4: "ARM", 0xaa64: "ARM64"}
                    out["pe_machine"] = machines.get(machine, hex(machine))
                    _num_sections = struct.unpack("<H", fh.read(2))[0]
                    _timestamp = struct.unpack("<I", fh.read(4))[0]
                    if _timestamp:
                        try:
                            out["pe_compiled"] = datetime.utcfromtimestamp(
                                _timestamp).isoformat(timespec="seconds")
                        except Exception:
                            pass
        except Exception:
            pass

    # Mach-O
    elif head[:4] in (b"\xfe\xed\xfa\xce", b"\xfe\xed\xfa\xcf",
                      b"\xce\xfa\xed\xfe", b"\xcf\xfa\xed\xfe"):
        out["format"] = "Mach-O"

    # shell / script heuristics
    else:
        first_line = head.split(b"\n", 1)[0]
        if first_line.startswith(b"#!"):
            out["format"] = "script"
            try:
                out["interpreter"] = first_line[2:].decode(
                    "utf-8", errors="ignore")[:200]
            except Exception:
                pass
        else:
            out["format"] = "unknown"

    # embedded strings (quick scan for interesting IOCs)
    try:
        with path.open("rb") as fh:
            raw = fh.read(1 << 20)
        text = raw.decode("latin-1", errors="ignore")
        urls = re.findall(r"https?://[^\s\"'<>]{6,200}", text)
        ips = re.findall(r"\b(?:\d{1,3}\.){3}\d{1,3}\b", text)
        emails = re.findall(
            r"\b[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}\b", text)
        if urls:
            out["strings_urls"] = list(dict.fromkeys(urls))[:25]
        if ips:
            out["strings_ips"] = list(dict.fromkeys(ips))[:25]
        if emails:
            out["strings_emails"] = list(dict.fromkeys(emails))[:25]
    except Exception:
        pass

    return out


# ============================================================
# Text
# ============================================================
def _text_meta(path: Path):
    out = {}
    try:
        with path.open("rb") as fh:
            raw = fh.read(1 << 20)
        try:
            text = raw.decode("utf-8")
            out["encoding"] = "utf-8"
        except UnicodeDecodeError:
            try:
                text = raw.decode("latin-1")
                out["encoding"] = "latin-1"
            except Exception:
                text = raw.decode("utf-8", errors="ignore")
                out["encoding"] = "utf-8 (lossy)"
        out["lines"] = text.count("\n") + 1
        out["words"] = len(text.split())
        out["chars"] = len(text)
        # BOM detection
        if raw.startswith(b"\xef\xbb\xbf"):
            out["bom"] = "utf-8-sig"
        elif raw.startswith(b"\xff\xfe"):
            out["bom"] = "utf-16-le"
        elif raw.startswith(b"\xfe\xff"):
            out["bom"] = "utf-16-be"
        # line-ending style
        if b"\r\n" in raw:
            out["line_ending"] = "crlf"
        elif b"\r" in raw:
            out["line_ending"] = "cr"
        elif b"\n" in raw:
            out["line_ending"] = "lf"
    except Exception as e:
        out["error"] = str(e)
    return out


# ============================================================
# Generic dispatch
# ============================================================
def extract_metadata(path: Path, deep=True, geocode=False, cfg=None):
    """Extract all metadata for a single file."""
    cfg = cfg or {}
    kind = classify(path)

    try:
        st = path.stat()
        base = {
            "path": str(path),
            "filename": path.name,
            "extension": path.suffix.lower(),
            "size_bytes": st.st_size,
            "kind": kind,
            "modified": datetime.fromtimestamp(st.st_mtime).isoformat(
                timespec="seconds"),
            "accessed": datetime.fromtimestamp(st.st_atime).isoformat(
                timespec="seconds"),
            "created_fs": datetime.fromtimestamp(st.st_ctime).isoformat(
                timespec="seconds"),
        }
    except Exception:
        base = {"path": str(path), "kind": kind}

    # always-on
    try:
        base.update(hash_file(path))
    except Exception:
        pass

    base["mime"] = detect_mime(path)
    desc = detect_description(path)
    if desc:
        base["file_description"] = desc

    ent = entropy_file(path)
    base["entropy"] = ent
    base["entropy_label"] = entropy_label(ent)

    # kind-specific
    try:
        if kind == "image":
            base.update(_image_meta(path, deep=deep))
        elif kind == "pdf":
            base.update(_pdf_meta(path))
        elif kind == "office":
            base.update(_office_meta(path))
        elif kind == "archive":
            base.update(_archive_meta(path))
        elif kind == "audio":
            base.update(_audio_meta(path))
        elif kind == "video":
            base.update(_video_meta(path))
        elif kind == "executable":
            base.update(_executable_meta(path))
        elif kind == "text":
            base.update(_text_meta(path))
    except Exception as e:
        base["extract_error"] = str(e)

    # reverse geocode GPS if requested
    if geocode and base.get("gps"):
        g = base["gps"]
        rev = reverse_geocode(g.get("lat"), g.get("lon"))
        if rev:
            base["gps"]["reverse_geocode"] = rev

    # timestamp anomaly detection
    base["timestamp_analysis"] = _timestamp_analysis(base)

    # sensitive-data scan
    base["sensitive_hits"] = _scan_sensitive(base)

    # risk score
    base["risk_score"], base["risk_reasons"] = _risk_score(base)

    return base


# ============================================================
# Timestamp anomaly
# ============================================================
def _parse_dt(v):
    if not v:
        return None
    if isinstance(v, datetime):
        return v
    for fmt in (
        "%Y:%m:%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%S.%f",
    ):
        try:
            return datetime.strptime(str(v)[:26].rstrip("Z"), fmt)
        except Exception:
            continue
    try:
        return datetime.fromisoformat(str(v).replace("Z", "+00:00"))
    except Exception:
        return None


def _timestamp_analysis(meta):
    """Compare fs mtime against EXIF/PDF/Office modified timestamps."""
    out = {"anomalies": [], "drift_days": {}}
    try:
        fs_mod = _parse_dt(meta.get("modified"))
    except Exception:
        fs_mod = None

    candidates = {}

    exif = meta.get("exif") or {}
    for key in ("DateTimeOriginal", "DateTime", "DateTimeDigitized"):
        dt = _parse_dt(exif.get(key))
        if dt:
            candidates[key] = dt

    for doc_key in ("metadata", "xmp", "docinfo", "docx_core", "xlsx_core",
                    "pptx_core", "docProps/core.xml"):
        block = meta.get(doc_key)
        if not isinstance(block, dict):
            continue
        for k in ("ModDate", "ModifyDate", "modified", "dcterms:modified",
                  "xmp:ModifyDate", "modified_by"):
            dt = _parse_dt(block.get(k))
            if dt:
                candidates[f"{doc_key}:{k}"] = dt

    if fs_mod:
        for name, dt in candidates.items():
            try:
                if dt.tzinfo is not None and fs_mod.tzinfo is None:
                    drift = (fs_mod - dt.replace(tzinfo=None)).days
                elif dt.tzinfo is None and fs_mod.tzinfo is not None:
                    drift = (fs_mod.replace(tzinfo=None) - dt).days
                else:
                    drift = (fs_mod - dt).days
            except Exception:
                continue
            out["drift_days"][name] = drift
            if abs(drift) > 365 * 5:
                out["anomalies"].append(
                    f"{name} differs from filesystem mtime by {drift} days")
            elif drift > 1 and name.lower().endswith(
                    ("modified", "moddate", "modifydate")):
                out["anomalies"].append(
                    f"metadata '{name}' predates filesystem mtime by "
                    f"{drift} day(s) — possible edit after copy")

    return out


# ============================================================
# Sensitive data scan
# ============================================================
SENSITIVE_PATTERNS = [
    ("email", r"\b[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}\b", 20),
    ("windows_path", r"[A-Z]:\\\\?[^\s\"'<>]{3,}", 15),
    ("unix_path", r"/(?:home|Users|root|etc|var|opt)/[^\s\"'<>]{2,}", 15),
    ("username", r"\b(?:author|creator|last[_\s]?modified[_\s]?by|owner)"
                 r"\s*[:=]\s*[A-Za-z0-9._\- ]{2,40}", 20),
    ("hostname", r"\b[a-z0-9\-]+\.(?:local|lan|internal|corp)\b", 25),
    ("ip_private", r"\b(?:10|172\.(?:1[6-9]|2\d|3[01])|192\.168)\.\d{1,3}\.\d{1,3}\b", 10),
    ("api_key", r"\b(?:api[_-]?key|token|secret)\s*[:=]\s*[^\s\"']{8,}", 45),
    ("password", r"(?:password|passwd|pwd)\s*[:=]\s*[^\s\"']{4,}", 45),
    ("jwt", r"\beyJ[A-Za-z0-9_\-]+\.[A-Za-z0-9_\-]+\.[A-Za-z0-9_\-]+\b", 40),
    ("aws_key", r"\bAKIA[0-9A-Z]{16}\b", 55),
    ("ssh_pubkey", r"ssh-rsa AAAA[0-9A-Za-z+/]+[=]{0,3}", 30),
    ("private_key", r"-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----", 65),
    ("credit_card", r"\b(?:\d{4}[-\s]?){3}\d{4}\b", 50),
    ("ssn_us", r"\b\d{3}-\d{2}-\d{4}\b", 45),
]


def _iter_strings(obj, depth=0, max_depth=6):
    """Recursively yield every string value in a nested structure."""
    if depth > max_depth:
        return
    if isinstance(obj, str):
        yield obj
    elif isinstance(obj, dict):
        for v in obj.values():
            yield from _iter_strings(v, depth + 1, max_depth)
    elif isinstance(obj, (list, tuple, set)):
        for v in obj:
            yield from _iter_strings(v, depth + 1, max_depth)


def _scan_sensitive(meta):
    """Regex scan of every string field, deduped with context."""
    hits = {}
    for s in _iter_strings(meta):
        if not s or len(s) > 8000:
            continue
        for name, pattern, weight in SENSITIVE_PATTERNS:
            try:
                for m in re.findall(pattern, s, re.IGNORECASE):
                    key = name
                    hits.setdefault(key, {
                        "weight": weight,
                        "matches": [],
                    })
                    if m not in hits[key]["matches"]:
                        hits[key]["matches"].append(m[:120])
                    if len(hits[key]["matches"]) >= 5:
                        break
            except Exception:
                continue
    # trim
    out = {}
    for k, v in hits.items():
        out[k] = {
            "weight": v["weight"],
            "count": len(v["matches"]),
            "samples": v["matches"][:3],
        }
    return out


# ============================================================
# Risk scoring
# ============================================================
def _risk_score(meta):
    score = 0
    reasons = []

    # GPS in image
    gps = meta.get("gps")
    if gps and gps.get("lat") is not None:
        score += 25
        reasons.append("GPS coordinates embedded in metadata")

    # author / creator fields
    exif = meta.get("exif") or {}
    if exif.get("Artist"):
        score += 10
        reasons.append(f"EXIF Artist = {str(exif['Artist'])[:60]}")
    if exif.get("Copyright"):
        score += 5
        reasons.append("EXIF Copyright present")
    if exif.get("Make") or exif.get("Model"):
        score += 5
        reasons.append(
            f"camera identified: {exif.get('Make', '?')} "
            f"{exif.get('Model', '?')}"
        )
    if exif.get("Software"):
        score += 5
        reasons.append(f"software tag: {str(exif['Software'])[:60]}")
    if exif.get("BodySerialNumber") or exif.get("CameraSerialNumber"):
        score += 15
        reasons.append("camera body serial number present")

    # sensitive hits
    sens = meta.get("sensitive_hits") or {}
    for name, info in sens.items():
        w = info.get("weight", 10)
        score = max(score, min(100, score + w))
        reasons.append(f"sensitive: {name} ({info.get('count', 0)} hit(s))")

    # high entropy — could be packed/encrypted
    ent = meta.get("entropy") or 0
    if ent > 7.9 and meta.get("kind") not in ("archive", "audio", "video",
                                              "image"):
        score += 10
        reasons.append(
            f"very high entropy ({ent}) — possibly packed or encrypted"
        )

    # timestamp anomalies
    tsa = meta.get("timestamp_analysis") or {}
    for anom in tsa.get("anomalies", []):
        score += 8
        reasons.append(anom)

    # stego suspicion
    stego = (meta.get("stego") or {})
    if stego.get("likelihood") == "high":
        score += 20
        reasons.append("high likelihood of LSB steganography")
    elif stego.get("likelihood") == "medium":
        score += 10
        reasons.append("possible LSB steganography")

    # embedded URLs / IOCs in executables
    if meta.get("strings_urls"):
        score += 8
        reasons.append(
            f"{len(meta['strings_urls'])} URL(s) in binary strings"
        )
    if meta.get("strings_ips"):
        score += 8
        reasons.append(
            f"{len(meta['strings_ips'])} IP(s) in binary strings"
        )

    return min(score, 100), reasons


def _risk_label(score):
    if score < 10:
        return "very low", "green"
    if score < 30:
        return "low", "green"
    if score < 55:
        return "moderate", "yellow"
    if score < 80:
        return "high", "red"
    return "critical", "red"


# ============================================================
# Scrubbers
# ============================================================
def scrub_image(path: Path):
    """Remove all EXIF / text chunks by re-encoding pixel data."""
    if not HAS_PIL:
        return None
    try:
        out_dir = OUT_DIR / "scrubbed"
        out_dir.mkdir(parents=True, exist_ok=True)
        out = out_dir / f"clean_{path.stem}{path.suffix}"
        img = Image.open(path)
        data = list(img.getdata())
        clean = Image.new(img.mode, img.size)
        clean.putdata(data)
        # save with no metadata
        ext = path.suffix.lower()
        fmt = "JPEG" if ext in (".jpg", ".jpeg") else \
              "PNG" if ext == ".png" else \
              "WEBP" if ext == ".webp" else img.format or "PNG"
        if fmt == "JPEG":
            clean.save(out, fmt, quality=95)
        else:
            clean.save(out, fmt)
        return out
    except Exception as e:
        console.print(f"[dg]  image scrub error: {e}[/dg]")
        return None


def scrub_pdf(path: Path):
    """Strip document info + XMP from a PDF using pikepdf."""
    if not HAS_PIKEPDF:
        return None
    try:
        out_dir = OUT_DIR / "scrubbed"
        out_dir.mkdir(parents=True, exist_ok=True)
        out = out_dir / f"clean_{path.stem}.pdf"
        with pikepdf.open(str(path)) as pdf:
            try:
                if pdf.docinfo:
                    for k in list(pdf.docinfo.keys()):
                        del pdf.docinfo[k]
            except Exception:
                pass
            try:
                with pdf.open_metadata() as meta:
                    for k in list(meta.keys()):
                        try:
                            del meta[k]
                        except Exception:
                            pass
            except Exception:
                pass
            pdf.save(str(out))
        return out
    except Exception as e:
        console.print(f"[dg]  pdf scrub error: {e}[/dg]")
        return None


def scrub_office(path: Path):
    """Remove docProps/*.xml from an Office zip container."""
    try:
        out_dir = OUT_DIR / "scrubbed"
        out_dir.mkdir(parents=True, exist_ok=True)
        out = out_dir / f"clean_{path.stem}{path.suffix}"
        with zipfile.ZipFile(path) as src, \
                zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():
                name = item.filename
                if name.startswith("docProps/"):
                    continue
                if name == "_rels/.rels":
                    raw = src.read(name).decode("utf-8", errors="ignore")
                    raw = re.sub(
                        r"<Relationship[^>]*docProps[^>]*/>", "", raw
                    )
                    dst.writestr(name, raw)
                    continue
                dst.writestr(item, src.read(name))
        return out
    except Exception as e:
        console.print(f"[dg]  office scrub error: {e}[/dg]")
        return None


def scrub_archive(path: Path):
    """Rebuild zip/tar without per-entry timestamps or comments."""
    try:
        out_dir = OUT_DIR / "scrubbed"
        out_dir.mkdir(parents=True, exist_ok=True)
        out = out_dir / f"clean_{path.stem}{path.suffix}"
        ext = path.suffix.lower()
        if ext in (".zip", ".jar", ".apk", ".epub"):
            with zipfile.ZipFile(path) as src, \
                    zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
                for item in src.infolist():
                    new = zipfile.ZipInfo(
                        item.filename, date_time=(1980, 1, 1, 0, 0, 0))
                    new.compress_type = item.compress_type
                    new.external_attr = item.external_attr
                    new.internal_attr = item.internal_attr
                    new.create_system = item.create_system
                    dst.writestr(new, src.read(item.filename))
            return out
        if ext in (".tar", ".gz", ".tgz", ".bz2", ".xz"):
            with tarfile.open(str(path), "r:*") as src, \
                    tarfile.open(str(out), "w") as dst:
                for member in src.getmembers():
                    member.mtime = 0
                    member.uid = 0
                    member.gid = 0
                    member.uname = ""
                    member.gname = ""
                    data = src.extractfile(member)
                    dst.addfile(member, data)
            return out
    except Exception as e:
        console.print(f"[dg]  archive scrub error: {e}[/dg]")
        return None


# ============================================================
# History
# ============================================================
def _hist_key(path: Path):
    try:
        st = path.stat()
        return hashlib.sha256(
            f"{path.resolve()}:{st.st_size}".encode()).hexdigest()[:24]
    except Exception:
        return hashlib.sha256(str(path).encode()).hexdigest()[:24]


def _load_hist(path: Path):
    p = HIST_DIR / f"{_hist_key(path)}.json"
    if p.exists():
        try:
            return json.loads(p.read_text(encoding="utf-8"))
        except Exception:
            return None
    return None


def _save_hist(path: Path, snapshot: dict):
    p = HIST_DIR / f"{_hist_key(path)}.json"
    try:
        p.write_text(json.dumps(snapshot, indent=2, ensure_ascii=False,
                                default=str), encoding="utf-8")
    except Exception:
        pass


def _diff_hist(prev, curr):
    if not prev:
        return []
    diffs = []
    if prev.get("sha256") and curr.get("sha256") and \
            prev["sha256"] != curr["sha256"]:
        diffs.append("file hash changed since last scan")
    if prev.get("size_bytes") and curr.get("size_bytes") and \
            prev["size_bytes"] != curr["size_bytes"]:
        diffs.append(
            f"size: {prev['size_bytes']} → {curr['size_bytes']}"
        )
    if prev.get("risk_score") != curr.get("risk_score"):
        diffs.append(
            f"risk score: {prev.get('risk_score')} → "
            f"{curr.get('risk_score')}"
        )
    return diffs


# ============================================================
# Render
# ============================================================
def _render(meta, history_diffs=None):
    name = Path(meta.get("path", "?")).name
    score = meta.get("risk_score", 0)
    label, color = _risk_label(score)

    console.print(Panel.fit(
        f"[bg]METADATA[/bg]  [dg]{name}[/dg]  ·  "
        f"[g]{meta.get('kind')}[/g]  ·  "
        f"[{color}]risk {score}/100[/{color}]",
        border_style="bd", padding=(0, 2),
    ))

    # file
    ft = Table(border_style="bd", title="[bg]file[/bg]",
               show_header=False, padding=(0, 2))
    ft.add_column("k", style="dg", no_wrap=True)
    ft.add_column("v", style="g", overflow="fold")
    for k in ("path", "mime", "file_description", "size_bytes", "modified",
              "accessed", "created_fs", "entropy", "entropy_label",
              "md5", "sha1", "sha256"):
        v = meta.get(k)
        if v is not None:
            ft.add_row(k, str(v))
    console.print(ft)

    # risk reasons
    reasons = meta.get("risk_reasons") or []
    if reasons:
        rt = Table(border_style="bd",
                   title=f"[bg]risk reasons[/bg]  "
                         f"[{color}]{label}[/{color}]",
                   show_header=False, padding=(0, 2))
        rt.add_column("r", style="g", overflow="fold")
        for r in reasons:
            rt.add_row(f"· {r}")
        console.print(rt)

    # kind-specific
    kind = meta.get("kind")
    if kind == "image":
        _render_image(meta)
    elif kind == "pdf":
        _render_pdf(meta)
    elif kind == "office":
        _render_office(meta)
    elif kind == "archive":
        _render_archive(meta)
    elif kind == "audio":
        _render_audio(meta)
    elif kind == "video":
        _render_video(meta)
    elif kind == "executable":
        _render_exe(meta)
    elif kind == "text":
        _render_text(meta)

    # timestamp analysis
    tsa = meta.get("timestamp_analysis") or {}
    if tsa.get("anomalies"):
        tt = Table(border_style="bd",
                   title="[bg]timestamp anomalies[/bg]",
                   show_header=False, padding=(0, 2))
        tt.add_column("a", style="g", overflow="fold")
        for a in tsa["anomalies"]:
            tt.add_row(f"· {a}")
        console.print(tt)

    # sensitive hits
    sens = meta.get("sensitive_hits") or {}
    if sens:
        st = Table(border_style="bd",
                   title=f"[bg]sensitive data[/bg]  "
                         f"[yellow]{len(sens)} type(s)[/yellow]",
                   show_header=True, padding=(0, 2))
        st.add_column("type", style="bg")
        st.add_column("count", style="dg", justify="right")
        st.add_column("samples", style="g", overflow="fold")
        for name, info in sorted(
            sens.items(), key=lambda x: -x[1].get("weight", 0)
        ):
            st.add_row(name, str(info.get("count", 0)),
                       "; ".join(info.get("samples", []))[:150])
        console.print(st)

    # history diff
    if history_diffs:
        console.print(Panel.fit(
            "[bg]changes since last scan[/bg]",
            border_style="bd", padding=(0, 2),
        ))
        for d in history_diffs:
            console.print(f"[yellow]  · {d}[/yellow]")
        console.print()

    console.print()


def _render_image(meta):
    info = {k: meta.get(k) for k in
            ("format", "mode", "width", "height", "aspect_ratio",
             "megapixels", "exif_tag_count", "blur_score", "dominant_hex",
             "brightness", "contrast")}
    info = {k: v for k, v in info.items() if v is not None}
    if info:
        it = Table(border_style="bd", title="[bg]image[/bg]",
                   show_header=False, padding=(0, 2))
        it.add_column("k", style="dg", no_wrap=True)
        it.add_column("v", style="g", overflow="fold")
        for k, v in info.items():
            it.add_row(k, str(v))
        console.print(it)

    gps = meta.get("gps")
    if gps:
        gt = Table(border_style="bd", title="[bg]gps[/bg]",
                   show_header=False, padding=(0, 2))
        gt.add_column("k", style="dg", no_wrap=True)
        gt.add_column("v", style="g", overflow="fold")
        for k in ("lat", "lon", "altitude_m", "gps_timestamp"):
            if gps.get(k) is not None:
                gt.add_row(k, str(gps[k]))
        if gps.get("osm"):
            gt.add_row("osm", gps["osm"])
        rev = gps.get("reverse_geocode") or {}
        if rev:
            gt.add_row("address", (rev.get("display_name") or "")[:150])
            addr = rev.get("address") or {}
            for key in ("country", "state", "city", "postcode", "road"):
                if addr.get(key):
                    gt.add_row(f"  {key}", str(addr[key]))
        console.print(gt)

    exif = meta.get("exif") or {}
    if exif:
        et = Table(border_style="bd",
                   title=f"[bg]exif ({len(exif)})[/bg]",
                   show_header=True, padding=(0, 2))
        et.add_column("tag", style="bg", overflow="fold")
        et.add_column("value", style="g", overflow="fold")
        for k, v in list(exif.items())[:50]:
            if k == "GPSInfo":
                continue
            et.add_row(str(k), str(v)[:140])
        console.print(et)

    stego = meta.get("stego")
    if stego:
        color = {"high": "red", "medium": "yellow",
                 "low": "green"}.get(stego.get("likelihood"), "dg")
        st = Table(border_style="bd",
                   title=f"[bg]stego[/bg]  "
                         f"[{color}]{stego.get('likelihood')}[/{color}]",
                   show_header=False, padding=(0, 2))
        st.add_column("k", style="dg", no_wrap=True)
        st.add_column("v", style="g", overflow="fold")
        for k in ("lsb_ratio_r", "lsb_ratio_g", "lsb_ratio_b",
                  "chi_square_avg"):
            if stego.get(k) is not None:
                st.add_row(k, str(stego[k]))
        if stego.get("suspicious"):
            st.add_row("suspicious", ", ".join(stego["suspicious"]))
        console.print(st)


def _render_pdf(meta):
    doc_keys = ("metadata", "xmp", "docinfo", "pikepdf_xmp")
    for key in doc_keys:
        block = meta.get(key)
        if isinstance(block, dict) and block:
            bt = Table(border_style="bd",
                       title=f"[bg]pdf {key} ({len(block)})[/bg]",
                       show_header=True, padding=(0, 2))
            bt.add_column("k", style="bg", overflow="fold")
            bt.add_column("v", style="g", overflow="fold")
            for k, v in list(block.items())[:50]:
                bt.add_row(str(k), str(v)[:140])
            console.print(bt)
    info = {}
    for k in ("pages", "pages_pikepdf", "pdf_version", "encrypted",
              "attachments"):
        if meta.get(k) is not None:
            info[k] = meta[k]
    if info:
        it = Table(border_style="bd", title="[bg]pdf structure[/bg]",
                   show_header=False, padding=(0, 2))
        it.add_column("k", style="dg", no_wrap=True)
        it.add_column("v", style="g", overflow="fold")
        for k, v in info.items():
            it.add_row(k, str(v))
        console.print(it)


def _render_office(meta):
    for key in ("docx_core", "xlsx_core", "pptx_core",
                "docProps/core.xml", "docProps/app.xml",
                "docProps/custom.xml", "meta.xml"):
        block = meta.get(key)
        if isinstance(block, dict) and any(block.values()):
            bt = Table(border_style="bd",
                       title=f"[bg]{key}[/bg]",
                       show_header=True, padding=(0, 2))
            bt.add_column("k", style="bg", overflow="fold")
            bt.add_column("v", style="g", overflow="fold")
            for k, v in list(block.items())[:40]:
                if v in (None, ""):
                    continue
                bt.add_row(str(k), str(v)[:140])
            console.print(bt)
    extras = {}
    for k in ("paragraphs", "tables", "slides", "sheets", "entries"):
        if meta.get(k) is not None:
            extras[k] = meta[k]
    if extras:
        it = Table(border_style="bd", title="[bg]structure[/bg]",
                   show_header=False, padding=(0, 2))
        it.add_column("k", style="dg", no_wrap=True)
        it.add_column("v", style="g", overflow="fold")
        for k, v in extras.items():
            it.add_row(k, str(v))
        console.print(it)


def _render_archive(meta):
    info = {}
    for k in ("archive_type", "entries", "total_size", "compressed_size",
              "comment"):
        if meta.get(k) is not None:
            info[k] = meta[k]
    if info:
        it = Table(border_style="bd", title="[bg]archive[/bg]",
                   show_header=False, padding=(0, 2))
        it.add_column("k", style="dg", no_wrap=True)
        it.add_column("v", style="g", overflow="fold")
        for k, v in info.items():
            it.add_row(k, str(v)[:200])
        console.print(it)
    sample = meta.get("sample") or []
    if sample:
        console.print("[dg]  entries (first 30):[/dg]")
        for s in sample:
            console.print(f"[g]    · {s}[/g]")
        console.print()


def _render_audio(meta):
    info = {k: meta.get(k) for k in
            ("mutagen_type", "length", "bitrate", "sample_rate",
             "channels", "bits_per_sample", "codec", "codec_description",
             "tag_count")}
    info = {k: v for k, v in info.items() if v is not None}
    if info:
        it = Table(border_style="bd", title="[bg]audio[/bg]",
                   show_header=False, padding=(0, 2))
        it.add_column("k", style="dg", no_wrap=True)
        it.add_column("v", style="g", overflow="fold")
        for k, v in info.items():
            it.add_row(k, str(v))
        console.print(it)
    tags = meta.get("tags") or {}
    if tags:
        tt = Table(border_style="bd",
                   title=f"[bg]audio tags ({len(tags)})[/bg]",
                   show_header=True, padding=(0, 2))
        tt.add_column("tag", style="bg", overflow="fold")
        tt.add_column("value", style="g", overflow="fold")
        for k, v in list(tags.items())[:40]:
            tt.add_row(str(k), str(v)[:140])
        console.print(tt)


def _render_video(meta):
    info = {k: meta.get(k) for k in
            ("format_name", "format_long_name", "duration", "size",
             "bit_rate")}
    info = {k: v for k, v in info.items() if v is not None}
    if info:
        it = Table(border_style="bd", title="[bg]video format[/bg]",
                   show_header=False, padding=(0, 2))
        it.add_column("k", style="dg", no_wrap=True)
        it.add_column("v", style="g", overflow="fold")
        for k, v in info.items():
            it.add_row(k, str(v))
        console.print(it)
    ftags = meta.get("format_tags") or {}
    if ftags:
        ft = Table(border_style="bd", title="[bg]format tags[/bg]",
                   show_header=True, padding=(0, 2))
        ft.add_column("tag", style="bg", overflow="fold")
        ft.add_column("value", style="g", overflow="fold")
        for k, v in list(ftags.items())[:40]:
            ft.add_row(str(k), str(v)[:140])
        console.print(ft)
    streams = meta.get("streams") or []
    if streams:
        st = Table(border_style="bd",
                   title=f"[bg]streams ({len(streams)})[/bg]",
                   show_header=True, padding=(0, 2))
        st.add_column("#", style="dg", justify="right")
        st.add_column("type", style="bg")
        st.add_column("codec", style="g")
        st.add_column("size", style="dg")
        st.add_column("details", style="dg", overflow="fold")
        for s in streams:
            details = []
            if s.get("width") and s.get("height"):
                details.append(f"{s['width']}x{s['height']}")
            if s.get("r_frame_rate"):
                details.append(f"fps {s['r_frame_rate']}")
            if s.get("channels"):
                details.append(f"ch {s['channels']}")
            if s.get("sample_rate"):
                details.append(f"{s['sample_rate']} Hz")
            st.add_row(
                str(s.get("index", "?")),
                s.get("codec_type") or "?",
                s.get("codec_name") or "?",
                str(s.get("bit_rate") or "—"),
                ", ".join(details)[:120],
            )
        console.print(st)


def _render_exe(meta):
    info = {k: meta.get(k) for k in
            ("format", "elf_class", "elf_endian", "elf_type", "elf_machine",
             "pe_machine", "pe_compiled", "interpreter")}
    info = {k: v for k, v in info.items() if v is not None}
    if info:
        it = Table(border_style="bd", title="[bg]executable[/bg]",
                   show_header=False, padding=(0, 2))
        it.add_column("k", style="dg", no_wrap=True)
        it.add_column("v", style="g", overflow="fold")
        for k, v in info.items():
            it.add_row(k, str(v))
        console.print(it)
    for key, title in (("strings_urls", "URLs"),
                       ("strings_ips", "IPs"),
                       ("strings_emails", "emails")):
        vals = meta.get(key) or []
        if vals:
            vt = Table(border_style="bd",
                       title=f"[bg]strings — {title} ({len(vals)})[/bg]",
                       show_header=False, padding=(0, 2))
            vt.add_column("v", style="g", overflow="fold")
            for s in vals[:25]:
                vt.add_row(str(s)[:160])
            console.print(vt)


def _render_text(meta):
    info = {k: meta.get(k) for k in
            ("encoding", "bom", "line_ending", "lines", "words", "chars")}
    info = {k: v for k, v in info.items() if v is not None}
    if info:
        it = Table(border_style="bd", title="[bg]text[/bg]",
                   show_header=False, padding=(0, 2))
        it.add_column("k", style="dg", no_wrap=True)
        it.add_column("v", style="g", overflow="fold")
        for k, v in info.items():
            it.add_row(k, str(v))
        console.print(it)


# ============================================================
# Exports
# ============================================================
def _export_json(all_meta, label="batch"):
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe = re.sub(r"[^A-Za-z0-9_.-]", "_", label)[:40]
    p = OUT_DIR / f"meta_{safe}_{ts}.json"
    p.write_text(json.dumps(all_meta, indent=2, ensure_ascii=False,
                            default=str), encoding="utf-8")
    console.print(f"[dg]  exported → {p}[/dg]")
    return p


def _export_csv(all_meta, label="batch"):
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe = re.sub(r"[^A-Za-z0-9_.-]", "_", label)[:40]
    p = OUT_DIR / f"meta_{safe}_{ts}.csv"
    with p.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["path", "kind", "size_bytes", "mime", "md5", "sha1",
                    "sha256", "entropy", "risk_score", "risk_reasons",
                    "has_gps", "gps_lat", "gps_lon",
                    "author", "software", "make", "model"])
        for m in all_meta:
            exif = m.get("exif") or {}
            gps = m.get("gps") or {}
            w.writerow([
                m.get("path"), m.get("kind"), m.get("size_bytes"),
                m.get("mime"), m.get("md5"), m.get("sha1"),
                m.get("sha256"), m.get("entropy"),
                m.get("risk_score"),
                "; ".join(m.get("risk_reasons") or []),
                bool(gps.get("lat") is not None),
                gps.get("lat") or "", gps.get("lon") or "",
                exif.get("Artist") or "", exif.get("Software") or "",
                exif.get("Make") or "", exif.get("Model") or "",
            ])
    console.print(f"[dg]  exported → {p}[/dg]")
    return p


def _export_html(all_meta, label="batch"):
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe = re.sub(r"[^A-Za-z0-9_.-]", "_", label)[:40]
    p = OUT_DIR / f"meta_{safe}_{ts}.html"

    def esc(x):
        return html.escape(str(x))

    rows_html = ""
    for m in all_meta:
        name = Path(m.get("path", "?")).name
        gps = m.get("gps") or {}
        map_link = gps.get("google") or ""
        score = m.get("risk_score", 0)
        _, color = _risk_label(score)
        color_hex = {"green": "#8fe388", "yellow": "#f2c14e",
                     "red": "#ff6b6b"}.get(color, "#ddd")
        rows_html += (
            "<div class='file'>"
            f"<h3>{esc(name)} <span style='color:{color_hex}'>"
            f"risk {score}</span></h3>"
            f"<p><b>path:</b> {esc(m.get('path'))}<br>"
            f"<b>kind:</b> {esc(m.get('kind'))} · "
            f"<b>size:</b> {esc(m.get('size_bytes'))} · "
            f"<b>mime:</b> {esc(m.get('mime'))}<br>"
            f"<b>md5:</b> {esc(m.get('md5'))}<br>"
            f"<b>sha256:</b> {esc(m.get('sha256'))}<br>"
            f"<b>entropy:</b> {esc(m.get('entropy'))} "
            f"({esc(m.get('entropy_label'))})</p>"
        )
        if map_link:
            rows_html += (
                f"<p><b>GPS:</b> {esc(gps.get('lat'))}, "
                f"{esc(gps.get('lon'))} — "
                f"<a href='{esc(map_link)}'>map</a></p>"
            )
        reasons = m.get("risk_reasons") or []
        if reasons:
            rows_html += "<ul>" + "".join(
                f"<li>{esc(r)}</li>" for r in reasons) + "</ul>"
        rows_html += "</div><hr>"

    doc = f"""<!doctype html>
<html><head><meta charset="utf-8">
<title>Metadata Report — {esc(label)}</title>
<style>
body {{ background:#111; color:#ddd; font-family:monospace; padding:2em; }}
h1, h2, h3 {{ color:#9fd4ff; }}
.file {{ margin:1em 0; }}
a {{ color:#8fe388; }}
</style></head><body>
<h1>Metadata Report — {esc(label)}</h1>
<p>{len(all_meta)} file(s) · generated
{esc(datetime.now().isoformat(timespec="seconds"))}</p>
{rows_html}
</body></html>
"""
    p.write_text(doc, encoding="utf-8")
    console.print(f"[dg]  exported → {p}[/dg]")
    return p


# ============================================================
# Batch processing
# ============================================================
def _gather_files(target: Path):
    if target.is_file():
        return [target]
    files = []
    for p in target.rglob("*"):
        if p.is_file() and classify(p) != "unknown":
            files.append(p)
    return files


def _process_batch(files, deep=True, geocode=False, cfg=None, workers=6):
    all_meta = []
    with tqdm(total=len(files), desc=f"parsing {len(files)} file(s)",
              ncols=60, bar_format=BAR_FMT) as bar:
        with ThreadPoolExecutor(max_workers=workers) as ex:
            futs = {
                ex.submit(extract_metadata, p, deep, geocode, cfg): p
                for p in files
            }
            for f in as_completed(futs):
                try:
                    all_meta.append(f.result())
                except Exception as e:
                    all_meta.append({
                        "path": str(futs[f]),
                        "error": str(e),
                    })
                bar.update(1)
    return all_meta


# ============================================================
# Runner
# ============================================================
def run():
    cfg = _load_config()

    while True:
        console.print(Panel.fit(
            "[bg]METADATA EXTRACTOR[/bg]\n"
            "[dg]file metadata forensics[/dg]",
            border_style="bd", padding=(0, 2),
        ))
        print("  1) extract metadata (file or directory)")
        print("  2) extract metadata with reverse geocoding")
        print("  3) deep mode (steganography + full analysis)")
        print("  4) batch directory scan + export")
        print("  5) scrub metadata from a file")
        print("  6) dependency check")
        print("  0) exit")
        print()

        choice = input("  meta >> ").strip().lower()

        if choice in ("0", "q", "exit", "quit"):
            return

        elif choice in ("1", "2", "3", "4"):
            raw = input("  path (file or directory) >> ").strip().strip('"')
            if not raw:
                console.print("[dg]  empty input[/dg]")
                continue
            target = Path(os.path.expanduser(raw))
            if not target.exists():
                console.print("[red]  [!] path does not exist[/red]")
                continue
            files = _gather_files(target)
            if not files:
                console.print("[yellow]  no supported files found[/yellow]")
                continue

            deep = choice == "3"
            geocode = choice in ("2", "3")
            workers = min(8, max(2, len(files)))

            all_meta = _process_batch(files, deep=deep, geocode=geocode,
                                      cfg=cfg, workers=workers)

            # render
            for m in all_meta:
                if "error" in m and not m.get("kind"):
                    console.print(f"[red]  ! {m.get('path')}: {m['error']}[/red]")
                    continue
                hd = []
                try:
                    prev = _load_hist(Path(m["path"]))
                    hd = _diff_hist(prev, m)
                    _save_hist(Path(m["path"]), {
                        "path": m.get("path"),
                        "sha256": m.get("sha256"),
                        "size_bytes": m.get("size_bytes"),
                        "risk_score": m.get("risk_score"),
                        "generated": datetime.now().isoformat(
                            timespec="seconds"),
                    })
                except Exception:
                    pass
                _render(m, hd)

            if choice == "4" or len(all_meta) > 1:
                if input("  export json? [y/N] >> ").strip().lower() == "y":
                    _export_json(all_meta, target.name or "batch")
                if input("  export csv? [y/N] >> ").strip().lower() == "y":
                    _export_csv(all_meta, target.name or "batch")
                if input("  export html? [y/N] >> ").strip().lower() == "y":
                    _export_html(all_meta, target.name or "batch")

        elif choice == "5":
            raw = input("  path to scrub >> ").strip().strip('"')
            if not raw:
                continue
            target = Path(os.path.expanduser(raw))
            if not target.is_file():
                console.print("[red]  [!] not a file[/red]")
                continue
            kind = classify(target)
            out = None
            if kind == "image":
                out = scrub_image(target)
            elif kind == "pdf":
                out = scrub_pdf(target)
            elif kind == "office":
                out = scrub_office(target)
            elif kind == "archive":
                out = scrub_archive(target)
            if out:
                console.print(f"[green]  scrubbed → {out}[/green]\n")
            else:
                console.print("[yellow]  no scrubber for this type "
                              "(or dependency missing)[/yellow]\n")

        elif choice == "6":
            deps = [
                ("Pillow", HAS_PIL),
                ("pypdf/PyPDF2", HAS_PYPDF),
                ("pikepdf", HAS_PIKEPDF),
                ("mutagen", HAS_MUTAGEN),
                ("python-magic", HAS_MAGIC),
                ("py7zr", HAS_PY7ZR),
                ("rarfile", HAS_RARFILE),
                ("python-docx", HAS_DOCX),
                ("openpyxl", HAS_OPENPYXL),
                ("python-pptx", HAS_PPTX),
            ]
            t = Table(border_style="bd", title="[bg]dependencies[/bg]",
                      show_header=True, padding=(0, 2))
            t.add_column("package", style="bg")
            t.add_column("status", style="g")
            for name, ok in deps:
                t.add_row(name, "[green]installed[/green]" if ok
                          else "[red]missing[/red]")
            # ffprobe
            ff = _ffprobe_available()
            t.add_row("ffprobe (video)", "[green]installed[/green]" if ff
                      else "[red]missing[/red]")
            console.print(t)
            console.print()
